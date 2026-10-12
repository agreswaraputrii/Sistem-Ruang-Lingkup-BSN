from pathlib import Path
from urllib.parse import urlencode
from io import BytesIO, StringIO

import csv


from fastapi import (
    APIRouter,
    Request,
    UploadFile,
    File,
    Form,
    HTTPException,
)


from fastapi.responses import (
    HTMLResponse,
    RedirectResponse,
    StreamingResponse,
)


from fastapi.templating import (
    Jinja2Templates,
)


from openpyxl import Workbook
from openpyxl.styles import Font


from app.database import get_connection


from app.services.certification_scheme_import import (
    import_certification_schemes,
)


# ============================================================
# ROUTER
# ============================================================

router = APIRouter()


# ============================================================
# TEMPLATE
# ============================================================

BASE_DIR = (
    Path(__file__)
    .resolve()
    .parent
    .parent
)


templates = Jinja2Templates(
    directory=str(
        BASE_DIR / "templates"
    )
)


# ============================================================
# COLUMNS
# ============================================================

CERTIFICATION_SCHEME_COLUMNS = [
    "id_skema_sertifikasi",
    "skema_sertifikasi",
]


# ============================================================
# FORMAT CODE
# ============================================================

def format_certification_scheme_code(
    number,
):

    return (
        f"SKM-{int(number):04d}"
    )


# ============================================================
# VALIDATE CODE
# ============================================================

def validate_certification_scheme_code(
    code,
):

    if not code:

        raise ValueError(
            "ID Skema Sertifikasi wajib diisi."
        )


    code_text = str(
        code
    )


    if not code_text.startswith(
        "SKM-"
    ):

        raise ValueError(
            (
                "Format ID Skema Sertifikasi tidak valid. "
                "Gunakan format SKM-XXXX, "
                "contoh SKM-0001."
            )
        )


    number_part = (
        code_text[4:]
    )


    if (
        len(number_part) != 4
        or
        not number_part.isdigit()
    ):

        raise ValueError(
            (
                "Format ID Skema Sertifikasi tidak valid. "
                "Gunakan format SKM-XXXX, "
                "contoh SKM-0001."
            )
        )


# ============================================================
# VALIDATE NAME
# ============================================================

def validate_certification_scheme_name(
    name,
):

    if not name:

        raise ValueError(
            "Nama Skema Sertifikasi wajib diisi."
        )


# ============================================================
# GENERATE NEXT CODE
# ============================================================

def generate_next_certification_scheme_code(
    conn,
):

    rows = (
        conn.execute(
            """
            SELECT
                certification_scheme_code

            FROM master_certification_scheme

            WHERE certification_scheme_code IS NOT NULL
            """
        )
        .fetchall()
    )


    max_number = 0


    for row in rows:

        code = (
            row[
                "certification_scheme_code"
            ]
        )


        if code is None:
            continue


        code_text = str(
            code
        )


        if not code_text.startswith(
            "SKM-"
        ):
            continue


        number_part = (
            code_text[4:]
        )


        if not number_part.isdigit():
            continue


        number = int(
            number_part
        )


        if number > max_number:

            max_number = number


    next_number = (
        max_number + 1
    )


    return (
        format_certification_scheme_code(
            next_number
        )
    )


# ============================================================
# FIND DUPLICATE CODE
# ============================================================

def find_duplicate_code(
    conn,
    code,
    exclude_id=None,
):

    query = """
        SELECT
            certification_scheme_id,
            certification_scheme_code,
            certification_scheme_name

        FROM master_certification_scheme

        WHERE certification_scheme_code = ?
    """


    params = [
        code,
    ]


    if exclude_id is not None:

        query += """
            AND certification_scheme_id != ?
        """

        params.append(
            exclude_id
        )


    query += """
        LIMIT 1
    """


    return (
        conn.execute(
            query,
            params,
        )
        .fetchone()
    )


# ============================================================
# FIND DUPLICATE NAME
# ============================================================

def find_duplicate_name(
    conn,
    name,
    exclude_id=None,
):

    query = """
        SELECT
            certification_scheme_id,
            certification_scheme_code,
            certification_scheme_name

        FROM master_certification_scheme

        WHERE certification_scheme_name = ?
    """


    params = [
        name,
    ]


    if exclude_id is not None:

        query += """
            AND certification_scheme_id != ?
        """

        params.append(
            exclude_id
        )


    query += """
        LIMIT 1
    """


    return (
        conn.execute(
            query,
            params,
        )
        .fetchone()
    )


# ============================================================
# LIST
# ============================================================

@router.get(
    "/certification-schemes",
    response_class=HTMLResponse,
)
def certification_scheme_list(
    request: Request,
    search: str = "",
    result: str = "",
    message: str = "",
):

    conn = get_connection()


    try:

        query = """
            SELECT
                certification_scheme_id,
                certification_scheme_code,
                certification_scheme_name

            FROM master_certification_scheme

            WHERE 1 = 1
        """


        params = []


        if search.strip():

            query += """
                AND (
                    LOWER(
                        certification_scheme_code
                    )
                    LIKE LOWER(?)

                    OR LOWER(
                        certification_scheme_name
                    )
                    LIKE LOWER(?)
                )
            """


            keyword = (
                f"%{search.strip()}%"
            )


            params.extend(
                [
                    keyword,
                    keyword,
                ]
            )


        query += """
            ORDER BY
                CAST(
                    SUBSTR(
                        certification_scheme_code,
                        5
                    )
                    AS INTEGER
                ),
                certification_scheme_name
        """


        items = (
            conn.execute(
                query,
                params,
            )
            .fetchall()
        )


        total = (
            conn.execute(
                """
                SELECT COUNT(*) AS total

                FROM master_certification_scheme
                """
            )
            .fetchone()["total"]
        )


    finally:

        conn.close()


    return templates.TemplateResponse(
        request=request,

        name="certification_schemes/list.html",

        context={
            "items":
                items,

            "search":
                search,

            "result":
                result,

            "message":
                message,

            "total":
                total,
        },
    )


# ============================================================
# CREATE PAGE
# ============================================================

@router.get(
    "/certification-schemes/create",
    response_class=HTMLResponse,
)
def certification_scheme_create_page(
    request: Request,
    status: str = "",
    message: str = "",
):

    conn = get_connection()


    try:

        next_code = (
            generate_next_certification_scheme_code(
                conn
            )
        )


    finally:

        conn.close()


    return templates.TemplateResponse(
        request=request,

        name="certification_schemes/create.html",

        context={
            "status":
                status,

            "message":
                message,

            "next_code":
                next_code,
        },
    )


# ============================================================
# CREATE PROCESS
# ============================================================

@router.post(
    "/certification-schemes/create"
)
def certification_scheme_create_process(
    certification_scheme_name: str = Form(...),
):

    conn = get_connection()


    try:

        validate_certification_scheme_name(
            certification_scheme_name
        )


        conn.execute(
            "BEGIN"
        )


        # ====================================================
        # GENERATE CODE FROM DATABASE
        # ====================================================

        certification_scheme_code = (
            generate_next_certification_scheme_code(
                conn
            )
        )


        validate_certification_scheme_code(
            certification_scheme_code
        )


        # ====================================================
        # DUPLICATE NAME
        # ====================================================

        duplicate_name = (
            find_duplicate_name(
                conn,
                certification_scheme_name,
            )
        )


        if duplicate_name:

            raise ValueError(
                (
                    "Nama Skema Sertifikasi "
                    f"'{certification_scheme_name}' "
                    "sudah tersedia dengan ID "
                    f"'{duplicate_name['certification_scheme_code']}'."
                )
            )


        # ====================================================
        # SAFETY DUPLICATE CODE
        # ====================================================

        duplicate_code = (
            find_duplicate_code(
                conn,
                certification_scheme_code,
            )
        )


        if duplicate_code:

            raise ValueError(
                (
                    "ID Skema Sertifikasi "
                    f"'{certification_scheme_code}' "
                    "sudah digunakan. "
                    "Silakan coba simpan kembali."
                )
            )


        # ====================================================
        # INSERT
        # ====================================================

        conn.execute(
            """
            INSERT INTO master_certification_scheme (
                certification_scheme_code,
                certification_scheme_name
            )

            VALUES (
                ?,
                ?
            )
            """,
            (
                certification_scheme_code,
                certification_scheme_name,
            ),
        )


        conn.commit()


    except Exception as exc:

        conn.rollback()


        query = urlencode(
            {
                "status":
                    "failed",

                "message":
                    str(exc),
            }
        )


        return RedirectResponse(
            url=(
                "/certification-schemes/create?"
                + query
            ),

            status_code=303,
        )


    finally:

        conn.close()


    query = urlencode(
        {
            "result":
                "success",

            "message":
                (
                    "Master Skema Sertifikasi "
                    f"berhasil ditambahkan dengan ID "
                    f"{certification_scheme_code}."
                ),
        }
    )


    return RedirectResponse(
        url=(
            "/certification-schemes?"
            + query
        ),

        status_code=303,
    )


# ============================================================
# EDIT PAGE
# ============================================================

@router.get(
    "/certification-schemes/{certification_scheme_id}/edit",
    response_class=HTMLResponse,
)
def certification_scheme_edit_page(
    request: Request,
    certification_scheme_id: int,
    status: str = "",
    message: str = "",
):

    conn = get_connection()


    try:

        item = (
            conn.execute(
                """
                SELECT
                    certification_scheme_id,
                    certification_scheme_code,
                    certification_scheme_name

                FROM master_certification_scheme

                WHERE certification_scheme_id = ?
                """,
                (
                    certification_scheme_id,
                ),
            )
            .fetchone()
        )


        if item is None:

            raise HTTPException(
                status_code=404,

                detail=(
                    "Master Skema Sertifikasi "
                    "tidak ditemukan."
                ),
            )


    finally:

        conn.close()


    return templates.TemplateResponse(
        request=request,

        name="certification_schemes/edit.html",

        context={
            "item":
                item,

            "status":
                status,

            "message":
                message,
        },
    )


# ============================================================
# EDIT PROCESS
# ============================================================

@router.post(
    "/certification-schemes/{certification_scheme_id}/edit"
)
def certification_scheme_edit_process(
    certification_scheme_id: int,
    certification_scheme_name: str = Form(...),
):

    conn = get_connection()


    try:

        existing = (
            conn.execute(
                """
                SELECT
                    certification_scheme_id,
                    certification_scheme_code,
                    certification_scheme_name

                FROM master_certification_scheme

                WHERE certification_scheme_id = ?
                """,
                (
                    certification_scheme_id,
                ),
            )
            .fetchone()
        )


        if existing is None:

            raise HTTPException(
                status_code=404,

                detail=(
                    "Master Skema Sertifikasi "
                    "tidak ditemukan."
                ),
            )


        validate_certification_scheme_name(
            certification_scheme_name
        )


        # ====================================================
        # DUPLICATE NAME
        # ====================================================

        duplicate_name = (
            find_duplicate_name(
                conn,
                certification_scheme_name,

                exclude_id=(
                    certification_scheme_id
                ),
            )
        )


        if duplicate_name:

            raise ValueError(
                (
                    "Nama Skema Sertifikasi "
                    f"'{certification_scheme_name}' "
                    "sudah tersedia dengan ID "
                    f"'{duplicate_name['certification_scheme_code']}'."
                )
            )


        # ====================================================
        # UPDATE NAME ONLY
        # ====================================================

        conn.execute(
            """
            UPDATE master_certification_scheme

            SET
                certification_scheme_name = ?,
                updated_at = CURRENT_TIMESTAMP

            WHERE certification_scheme_id = ?
            """,
            (
                certification_scheme_name,
                certification_scheme_id,
            ),
        )


        conn.commit()


    except HTTPException:

        conn.rollback()

        raise


    except Exception as exc:

        conn.rollback()


        query = urlencode(
            {
                "status":
                    "failed",

                "message":
                    str(exc),
            }
        )


        return RedirectResponse(
            url=(
                f"/certification-schemes/"
                f"{certification_scheme_id}/edit?"
                + query
            ),

            status_code=303,
        )


    finally:

        conn.close()


    query = urlencode(
        {
            "result":
                "success",

            "message":
                (
                    "Master Skema Sertifikasi "
                    "berhasil diperbarui."
                ),
        }
    )


    return RedirectResponse(
        url=(
            "/certification-schemes?"
            + query
        ),

        status_code=303,
    )


# ============================================================
# DELETE
# ============================================================

@router.post(
    "/certification-schemes/{certification_scheme_id}/delete"
)
def certification_scheme_delete(
    certification_scheme_id: int,
):

    conn = get_connection()


    try:

        existing = (
            conn.execute(
                """
                SELECT
                    certification_scheme_id,
                    certification_scheme_code,
                    certification_scheme_name

                FROM master_certification_scheme

                WHERE certification_scheme_id = ?
                """,
                (
                    certification_scheme_id,
                ),
            )
            .fetchone()
        )


        if existing is None:

            query = urlencode(
                {
                    "result":
                        "failed",

                    "message":
                        (
                            "Master Skema Sertifikasi "
                            "tidak ditemukan."
                        ),
                }
            )


            return RedirectResponse(
                url=(
                    "/certification-schemes?"
                    + query
                ),

                status_code=303,
            )


        conn.execute(
            """
            DELETE FROM master_certification_scheme

            WHERE certification_scheme_id = ?
            """,
            (
                certification_scheme_id,
            ),
        )


        conn.commit()


    except Exception as exc:

        conn.rollback()


        query = urlencode(
            {
                "result":
                    "failed",

                "message":
                    str(exc),
            }
        )


        return RedirectResponse(
            url=(
                "/certification-schemes?"
                + query
            ),

            status_code=303,
        )


    finally:

        conn.close()


    query = urlencode(
        {
            "result":
                "success",

            "message":
                (
                    "Master Skema Sertifikasi "
                    "berhasil dihapus."
                ),
        }
    )


    return RedirectResponse(
        url=(
            "/certification-schemes?"
            + query
        ),

        status_code=303,
    )


# ============================================================
# IMPORT PAGE
# ============================================================

@router.get(
    "/certification-schemes/import",
    response_class=HTMLResponse,
)
def certification_scheme_import_page(
    request: Request,
    status: str = "",
    message: str = "",
    import_mode: str = "",
    total_rows: int = 0,
    inserted_count: int = 0,
    existing_count: int = 0,
):

    import_result = None

    success_message = None

    error_message = None


    if status == "success":

        success_message = (
            message
            or
            "Master Skema Sertifikasi berhasil di-import."
        )


        import_result = {
            "import_mode":
                import_mode,

            "total_rows":
                total_rows,

            "inserted_count":
                inserted_count,

            "existing_count":
                existing_count,
        }


    elif status == "failed":

        error_message = (
            message
            or
            "Import Master Skema Sertifikasi gagal."
        )


    return templates.TemplateResponse(
        request=request,

        name="certification_schemes/import.html",

        context={
            "success_message":
                success_message,

            "error_message":
                error_message,

            "import_result":
                import_result,
        },
    )


# ============================================================
# IMPORT PROCESS
# ============================================================

@router.post(
    "/certification-schemes/import"
)
async def certification_scheme_import_process(
    file: UploadFile = File(...),
    import_mode: str = Form(...),
):

    try:

        if not file.filename:

            raise ValueError(
                "Nama file tidak ditemukan."
            )


        if not file.filename.lower().endswith(
            (
                ".xlsx",
                ".xls",
                ".csv",
            )
        ):

            raise ValueError(
                (
                    "Format file tidak didukung. "
                    "Gunakan .xlsx, .xls, atau .csv."
                )
            )


        file_content = (
            await file.read()
        )


        if not file_content:

            raise ValueError(
                "File yang di-upload kosong."
            )


        result = (
            import_certification_schemes(
                file_content=file_content,
                filename=file.filename,
                import_mode=import_mode,
            )
        )


        mode_name = (
            "Upload Ulang Semua Data"

            if import_mode == "replace"

            else "Tambah Data"
        )


        query = urlencode(
            {
                "status":
                    "success",

                "message":
                    (
                        "Import Master Skema Sertifikasi "
                        "berhasil dengan mode "
                        f"{mode_name}."
                    ),

                "import_mode":
                    result["import_mode"],

                "total_rows":
                    result["total_rows"],

                "inserted_count":
                    result["inserted_count"],

                "existing_count":
                    result["existing_count"],
            }
        )


        return RedirectResponse(
            url=(
                "/certification-schemes/import?"
                + query
            ),

            status_code=303,
        )


    except Exception as exc:

        query = urlencode(
            {
                "status":
                    "failed",

                "message":
                    str(exc),
            }
        )


        return RedirectResponse(
            url=(
                "/certification-schemes/import?"
                + query
            ),

            status_code=303,
        )


# ============================================================
# DOWNLOAD TEMPLATE
# ============================================================

@router.get(
    "/certification-schemes/template"
)
def download_certification_scheme_template():

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = (
        "Master Skema Sertifikasi"
    )


    for index, column in enumerate(
        CERTIFICATION_SCHEME_COLUMNS,
        start=1,
    ):

        cell = worksheet.cell(
            row=1,
            column=index,
            value=column,
        )

        cell.font = Font(
            bold=True
        )


    worksheet.column_dimensions[
        "A"
    ].width = 25


    worksheet.column_dimensions[
        "B"
    ].width = 60


    worksheet.freeze_panes = (
        "A2"
    )


    output = BytesIO()

    workbook.save(
        output
    )

    output.seek(0)


    return StreamingResponse(
        output,

        media_type=(
            "application/"
            "vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),

        headers={
            "Content-Disposition":
                (
                    "attachment; "
                    "filename="
                    "Template_Master_Skema_Sertifikasi.xlsx"
                )
        },
    )


# ============================================================
# EXPORT EXCEL
# ============================================================

@router.get(
    "/certification-schemes/export/xlsx"
)
def export_certification_scheme_xlsx():

    conn = get_connection()


    try:

        rows = (
            conn.execute(
                """
                SELECT
                    certification_scheme_code,
                    certification_scheme_name

                FROM master_certification_scheme

                ORDER BY
                    CAST(
                        SUBSTR(
                            certification_scheme_code,
                            5
                        )
                        AS INTEGER
                    ),
                    certification_scheme_name
                """
            )
            .fetchall()
        )


    finally:

        conn.close()


    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = (
        "Master Skema Sertifikasi"
    )


    worksheet.append(
        CERTIFICATION_SCHEME_COLUMNS
    )


    for cell in worksheet[1]:

        cell.font = Font(
            bold=True
        )


    for row in rows:

        worksheet.append(
            [
                row[
                    "certification_scheme_code"
                ],

                row[
                    "certification_scheme_name"
                ],
            ]
        )


    worksheet.column_dimensions[
        "A"
    ].width = 25


    worksheet.column_dimensions[
        "B"
    ].width = 60


    worksheet.freeze_panes = (
        "A2"
    )


    output = BytesIO()

    workbook.save(
        output
    )

    output.seek(0)


    return StreamingResponse(
        output,

        media_type=(
            "application/"
            "vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),

        headers={
            "Content-Disposition":
                (
                    "attachment; "
                    "filename="
                    "Master_Skema_Sertifikasi.xlsx"
                )
        },
    )


# ============================================================
# EXPORT CSV
# ============================================================

@router.get(
    "/certification-schemes/export/csv"
)
def export_certification_scheme_csv():

    conn = get_connection()


    try:

        rows = (
            conn.execute(
                """
                SELECT
                    certification_scheme_code,
                    certification_scheme_name

                FROM master_certification_scheme

                ORDER BY
                    CAST(
                        SUBSTR(
                            certification_scheme_code,
                            5
                        )
                        AS INTEGER
                    ),
                    certification_scheme_name
                """
            )
            .fetchall()
        )


    finally:

        conn.close()


    output = StringIO(
        newline=""
    )


    writer = csv.writer(
        output
    )


    writer.writerow(
        CERTIFICATION_SCHEME_COLUMNS
    )


    for row in rows:

        writer.writerow(
            [
                row[
                    "certification_scheme_code"
                ],

                row[
                    "certification_scheme_name"
                ],
            ]
        )


    csv_content = (
        "\ufeff"
        + output.getvalue()
    )


    return StreamingResponse(
        iter(
            [
                csv_content.encode(
                    "utf-8"
                )
            ]
        ),

        media_type=(
            "text/csv; charset=utf-8"
        ),

        headers={
            "Content-Disposition":
                (
                    "attachment; "
                    "filename="
                    "Master_Skema_Sertifikasi.csv"
                )
        },
    )