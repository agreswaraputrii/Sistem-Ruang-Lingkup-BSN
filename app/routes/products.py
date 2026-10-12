from pathlib import Path
from urllib.parse import urlencode
from io import BytesIO, StringIO

import csv
import re


from fastapi import (
    APIRouter,
    Request,
    UploadFile,
    File,
    Form,
    HTTPException,
)


from fastapi.responses import (
    HTMLResponse,
    RedirectResponse,
    StreamingResponse,
)


from fastapi.templating import (
    Jinja2Templates,
)


from openpyxl import Workbook

from openpyxl.styles import (
    Font,
    Alignment,
)


from app.database import (
    get_connection,
)


from app.services.product_import import (
    import_products,
)

from app.services.scope_batch_import import (
    recalculate_scope_batch,
)


# ============================================================
# ROUTER
# ============================================================

router = APIRouter()


# ============================================================
# TEMPLATE DIRECTORY
# ============================================================

BASE_DIR = (
    Path(__file__)
    .resolve()
    .parent
    .parent
)


templates = Jinja2Templates(
    directory=str(
        BASE_DIR
        / "templates"
    )
)


# ============================================================
# FORMAT MASTER PRODUK
# ============================================================

PRODUCT_COLUMNS = [
    "skema",
    "id_skema",
    "kategori",
    "id_kategori",
    "sub_kategori",
    "id_sub_kategori",
    "jenis",
    "id_jenis",
]


# ============================================================
# STATUS
# ============================================================

STATUS_VALUE = "VALUE"

STATUS_NOT_APPLICABLE = (
    "NOT_APPLICABLE"
)

STATUS_MISSING = "MISSING"


VALID_STATUSES = {
    STATUS_VALUE,
    STATUS_NOT_APPLICABLE,
    STATUS_MISSING,
}


# ============================================================
# HELPER EXPORT VALUE
# ============================================================

def export_value(
    status,
    value,
):

    if status == STATUS_NOT_APPLICABLE:
        return "-"

    if status == STATUS_MISSING:
        return ""

    return (
        value
        if value is not None
        else ""
    )


# ============================================================
# GET EXPORT DATA
# ============================================================

def get_product_export_data():

    conn = get_connection()

    try:

        rows = (
            conn.execute(
                """
                SELECT
                    path.path_id,

                    scheme.scheme_code,
                    scheme.scheme_name,

                    category.category_code,
                    category.category_name,

                    subcategory.subcategory_code,
                    subcategory.subcategory_name,

                    product.product_service_code,
                    product.product_service_name,

                    path.category_status,
                    path.subcategory_status,
                    path.product_service_status

                FROM master_product_path AS path

                JOIN master_scheme AS scheme
                    ON scheme.scheme_id
                    =
                    path.scheme_id

                LEFT JOIN master_category AS category
                    ON category.category_id
                    =
                    path.category_id

                LEFT JOIN master_subcategory AS subcategory
                    ON subcategory.subcategory_id
                    =
                    path.subcategory_id

                LEFT JOIN master_product_service AS product
                    ON product.product_service_id
                    =
                    path.product_service_id

                ORDER BY
                    scheme.scheme_code,
                    category.category_code,
                    subcategory.subcategory_code,
                    product.product_service_code,
                    path.path_id
                """
            )
            .fetchall()
        )


        export_rows = []


        for item in rows:

            export_rows.append(
                {
                    "skema":
                        item["scheme_name"],

                    "id_skema":
                        item["scheme_code"],

                    "kategori":
                        export_value(
                            item["category_status"],
                            item["category_name"],
                        ),

                    "id_kategori":
                        export_value(
                            item["category_status"],
                            item["category_code"],
                        ),

                    "sub_kategori":
                        export_value(
                            item["subcategory_status"],
                            item["subcategory_name"],
                        ),

                    "id_sub_kategori":
                        export_value(
                            item["subcategory_status"],
                            item["subcategory_code"],
                        ),

                    "jenis":
                        export_value(
                            item["product_service_status"],
                            item["product_service_name"],
                        ),

                    "id_jenis":
                        export_value(
                            item["product_service_status"],
                            item["product_service_code"],
                        ),
                }
            )


        return export_rows


    finally:

        conn.close()


# ============================================================
# GET OR CREATE SCHEME
# ============================================================

def get_or_create_scheme(
    conn,
    code,
    name,
):

    existing = (
        conn.execute(
            """
            SELECT scheme_id

            FROM master_scheme

            WHERE scheme_code = ?
            AND scheme_name = ?
            """,
            (
                code,
                name,
            ),
        )
        .fetchone()
    )


    if existing:
        return existing["scheme_id"]


    cursor = conn.execute(
        """
        INSERT INTO master_scheme (
            scheme_code,
            scheme_name
        )

        VALUES (
            ?,
            ?
        )
        """,
        (
            code,
            name,
        ),
    )


    return cursor.lastrowid


# ============================================================
# GET OR CREATE CATEGORY
# ============================================================

def get_or_create_category(
    conn,
    scheme_id,
    code,
    name,
):

    existing = (
        conn.execute(
            """
            SELECT category_id

            FROM master_category

            WHERE scheme_id = ?
            AND category_code = ?
            AND category_name = ?
            """,
            (
                scheme_id,
                code,
                name,
            ),
        )
        .fetchone()
    )


    if existing:
        return existing["category_id"]


    cursor = conn.execute(
        """
        INSERT INTO master_category (
            scheme_id,
            category_code,
            category_name
        )

        VALUES (
            ?,
            ?,
            ?
        )
        """,
        (
            scheme_id,
            code,
            name,
        ),
    )


    return cursor.lastrowid


# ============================================================
# GET OR CREATE SUBCATEGORY
# ============================================================

def get_or_create_subcategory(
    conn,
    scheme_id,
    category_id,
    code,
    name,
):

    if category_id is None:

        existing = (
            conn.execute(
                """
                SELECT subcategory_id

                FROM master_subcategory

                WHERE scheme_id = ?
                AND category_id IS NULL
                AND subcategory_code = ?
                AND subcategory_name = ?
                """,
                (
                    scheme_id,
                    code,
                    name,
                ),
            )
            .fetchone()
        )

    else:

        existing = (
            conn.execute(
                """
                SELECT subcategory_id

                FROM master_subcategory

                WHERE scheme_id = ?
                AND category_id = ?
                AND subcategory_code = ?
                AND subcategory_name = ?
                """,
                (
                    scheme_id,
                    category_id,
                    code,
                    name,
                ),
            )
            .fetchone()
        )


    if existing:
        return existing["subcategory_id"]


    cursor = conn.execute(
        """
        INSERT INTO master_subcategory (
            scheme_id,
            category_id,
            subcategory_code,
            subcategory_name
        )

        VALUES (
            ?,
            ?,
            ?,
            ?
        )
        """,
        (
            scheme_id,
            category_id,
            code,
            name,
        ),
    )


    return cursor.lastrowid


# ============================================================
# GET OR CREATE PRODUCT SERVICE
# ============================================================

def get_or_create_product_service(
    conn,
    scheme_id,
    category_id,
    subcategory_id,
    code,
    name,
):

    query = """
        SELECT product_service_id

        FROM master_product_service

        WHERE scheme_id = ?
        AND product_service_code = ?
        AND product_service_name = ?
    """


    params = [
        scheme_id,
        code,
        name,
    ]


    if category_id is None:

        query += """
            AND category_id IS NULL
        """

    else:

        query += """
            AND category_id = ?
        """

        params.append(
            category_id
        )


    if subcategory_id is None:

        query += """
            AND subcategory_id IS NULL
        """

    else:

        query += """
            AND subcategory_id = ?
        """

        params.append(
            subcategory_id
        )


    existing = (
        conn.execute(
            query,
            params,
        )
        .fetchone()
    )


    if existing:
        return existing["product_service_id"]


    cursor = conn.execute(
        """
        INSERT INTO master_product_service (
            scheme_id,
            category_id,
            subcategory_id,
            product_service_code,
            product_service_name
        )

        VALUES (
            ?,
            ?,
            ?,
            ?,
            ?
        )
        """,
        (
            scheme_id,
            category_id,
            subcategory_id,
            code,
            name,
        ),
    )


    return cursor.lastrowid


# ============================================================
# CHECK DUPLICATE ID JENIS
# ============================================================

def find_duplicate_product_code(
    conn,
    scheme_id,
    category_id,
    subcategory_id,
    product_service_code,
    exclude_path_id=None,
):

    query = """
        SELECT
            path.path_id,

            product.product_service_code,
            product.product_service_name

        FROM master_product_path AS path

        JOIN master_product_service AS product
            ON product.product_service_id
            =
            path.product_service_id

        WHERE path.scheme_id = ?

        AND path.product_service_status = 'VALUE'

        AND product.product_service_code = ?
    """


    params = [
        scheme_id,
        product_service_code,
    ]


    if category_id is None:

        query += """
            AND path.category_id IS NULL
        """

    else:

        query += """
            AND path.category_id = ?
        """

        params.append(
            category_id
        )


    if subcategory_id is None:

        query += """
            AND path.subcategory_id IS NULL
        """

    else:

        query += """
            AND path.subcategory_id = ?
        """

        params.append(
            subcategory_id
        )


    if exclude_path_id is not None:

        query += """
            AND path.path_id != ?
        """

        params.append(
            exclude_path_id
        )


    query += """
        LIMIT 1
    """


    return (
        conn.execute(
            query,
            params,
        )
        .fetchone()
    )


# ============================================================
# VALIDATE FORM
# ============================================================

def validate_product_form(
    scheme_name,
    scheme_code,

    category_status,
    category_name,
    category_code,

    subcategory_status,
    subcategory_name,
    subcategory_code,

    product_service_status,
    product_service_name,
    product_service_code,
):

    if not scheme_name:

        raise ValueError(
            "Nama Skema wajib diisi."
        )


    if not scheme_code:

        raise ValueError(
            "ID Skema wajib diisi."
        )


    if category_status not in VALID_STATUSES:

        raise ValueError(
            "Status Kategori tidak valid."
        )


    if subcategory_status not in VALID_STATUSES:

        raise ValueError(
            "Status Sub Kategori tidak valid."
        )


    if product_service_status not in VALID_STATUSES:

        raise ValueError(
            "Status Jenis tidak valid."
        )


    if category_status == STATUS_VALUE:

        if not category_name:

            raise ValueError(
                "Nama Kategori wajib diisi."
            )

        if not category_code:

            raise ValueError(
                "ID Kategori wajib diisi."
            )


    if subcategory_status == STATUS_VALUE:

        if not subcategory_name:

            raise ValueError(
                "Nama Sub Kategori wajib diisi."
            )

        if not subcategory_code:

            raise ValueError(
                "ID Sub Kategori wajib diisi."
            )


    if product_service_status == STATUS_VALUE:

        if not product_service_name:

            raise ValueError(
                "Nama Jenis wajib diisi."
            )

        if not product_service_code:

            raise ValueError(
                "ID Jenis wajib diisi."
            )


# ============================================================
# LOAD CREATE OPTIONS
# ============================================================

def get_product_form_options():

    conn = get_connection()


    try:

        schemes = [
            dict(row)

            for row in conn.execute(
                """
                SELECT
                    scheme_id,
                    scheme_code,
                    scheme_name

                FROM master_scheme

                ORDER BY
                    scheme_code,
                    scheme_name
                """
            )
            .fetchall()
        ]


        categories = [
            dict(row)

            for row in conn.execute(
                """
                SELECT
                    category_id,
                    scheme_id,
                    category_code,
                    category_name

                FROM master_category

                ORDER BY
                    scheme_id,
                    category_code,
                    category_name
                """
            )
            .fetchall()
        ]


        subcategories = [
            dict(row)

            for row in conn.execute(
                """
                SELECT
                    subcategory_id,
                    scheme_id,
                    category_id,
                    subcategory_code,
                    subcategory_name

                FROM master_subcategory

                ORDER BY
                    scheme_id,
                    category_id,
                    subcategory_code,
                    subcategory_name
                """
            )
            .fetchall()
        ]


        products = [
            dict(row)

            for row in conn.execute(
                """
                SELECT
                    product_service_id,
                    scheme_id,
                    category_id,
                    subcategory_id,
                    product_service_code,
                    product_service_name

                FROM master_product_service

                ORDER BY
                    scheme_id,
                    category_id,
                    subcategory_id,
                    product_service_code,
                    product_service_name
                """
            )
            .fetchall()
        ]


        return {
            "schemes":
                schemes,

            "categories":
                categories,

            "subcategories":
                subcategories,

            "product_services":
                products,
        }


    finally:

        conn.close()


# ============================================================
# LIST MASTER PRODUK
# ============================================================

@router.get(
    "/products",
    response_class=HTMLResponse,
)
def product_list(
    request: Request,
    search: str = "",
    result: str = "",
    message: str = "",
):

    conn = get_connection()


    try:

        query = """
            SELECT
                path.path_id,

                scheme.scheme_code,
                scheme.scheme_name,

                category.category_code,
                category.category_name,

                subcategory.subcategory_code,
                subcategory.subcategory_name,

                product.product_service_code,
                product.product_service_name,

                path.category_status,
                path.subcategory_status,
                path.product_service_status

            FROM master_product_path AS path

            JOIN master_scheme AS scheme
                ON scheme.scheme_id
                =
                path.scheme_id

            LEFT JOIN master_category AS category
                ON category.category_id
                =
                path.category_id

            LEFT JOIN master_subcategory AS subcategory
                ON subcategory.subcategory_id
                =
                path.subcategory_id

            LEFT JOIN master_product_service AS product
                ON product.product_service_id
                =
                path.product_service_id

            WHERE 1 = 1
        """


        params = []


        if search.strip():

            query += """
                AND (
                    LOWER(scheme.scheme_name)
                        LIKE LOWER(?)

                    OR LOWER(scheme.scheme_code)
                        LIKE LOWER(?)

                    OR LOWER(
                        COALESCE(
                            category.category_name,
                            ''
                        )
                    )
                        LIKE LOWER(?)

                    OR LOWER(
                        COALESCE(
                            category.category_code,
                            ''
                        )
                    )
                        LIKE LOWER(?)

                    OR LOWER(
                        COALESCE(
                            subcategory.subcategory_name,
                            ''
                        )
                    )
                        LIKE LOWER(?)

                    OR LOWER(
                        COALESCE(
                            subcategory.subcategory_code,
                            ''
                        )
                    )
                        LIKE LOWER(?)

                    OR LOWER(
                        COALESCE(
                            product.product_service_name,
                            ''
                        )
                    )
                        LIKE LOWER(?)

                    OR LOWER(
                        COALESCE(
                            product.product_service_code,
                            ''
                        )
                    )
                        LIKE LOWER(?)
                )
            """


            keyword = (
                f"%{search.strip()}%"
            )


            params.extend(
                [
                    keyword,
                    keyword,
                    keyword,
                    keyword,
                    keyword,
                    keyword,
                    keyword,
                    keyword,
                ]
            )


        query += """
            ORDER BY
                scheme.scheme_code,
                category.category_code,
                subcategory.subcategory_code,
                product.product_service_code,
                path.path_id
        """


        products = (
            conn.execute(
                query,
                params,
            )
            .fetchall()
        )


        total_paths = (
            conn.execute(
                """
                SELECT COUNT(*) AS total

                FROM master_product_path
                """
            )
            .fetchone()["total"]
        )


    finally:

        conn.close()


    return templates.TemplateResponse(
        request=request,

        name="products/list.html",

        context={
            "products":
                products,

            "search":
                search,

            "result":
                result,

            "message":
                message,

            "total_paths":
                total_paths,
        },
    )



# ============================================================
# REVIEW PREFILL HELPERS
# ============================================================

def normalize_prefill_text(
    value,
):
    if value is None:
        return ""

    return (
        " ".join(
            str(value)
            .replace("\n", " ")
            .replace("\r", " ")
            .replace("\t", " ")
            .split()
        )
        .strip()
    )


def prefill_status(
    value,
):
    text = normalize_prefill_text(
        value
    )

    if text == "-":
        return STATUS_NOT_APPLICABLE

    if not text:
        return STATUS_MISSING

    return STATUS_VALUE


def extract_prefill_code(
    value,
):
    """
    Mengambil local code terakhir dari tanda kurung.

    Contoh:
    'Peralatan kesehatan (01)' -> '01'
    'Peralatan kesehatan pendukung (01.03)' -> '03'
    'Jasa boga (05.02)' -> '02'
    """
    text = normalize_prefill_text(
        value
    )

    matches = re.findall(
        r"\(([^()]*)\)\s*$",
        text,
    )

    if not matches:
        return ""

    code_text = (
        matches[-1]
        .strip()
    )

    parts = [
        part.strip()
        for part
        in code_text.split(".")
        if part.strip()
    ]

    if not parts:
        return ""

    return parts[-1]


def remove_prefill_code(
    value,
):
    text = normalize_prefill_text(
        value
    )

    return (
        re.sub(
            r"\s*\([^()]*\)\s*$",
            "",
            text,
        )
        .strip()
    )


def find_prefill_scheme(
    schemes,
    source_value,
):
    source_name = (
        remove_prefill_code(
            source_value
        )
        .casefold()
    )

    if not source_name:
        return None

    for item in schemes:
        if (
            normalize_prefill_text(
                item[
                    "scheme_name"
                ]
            )
            .casefold()
            ==
            source_name
        ):
            return item

    return None


def find_prefill_category(
    categories,
    scheme_id,
    source_value,
):
    if (
        prefill_status(
            source_value
        )
        !=
        STATUS_VALUE
    ):
        return None

    source_code = (
        extract_prefill_code(
            source_value
        )
    )

    source_name = (
        remove_prefill_code(
            source_value
        )
        .casefold()
    )

    candidates = [
        item
        for item
        in categories
        if item[
            "scheme_id"
        ]
        ==
        scheme_id
    ]

    if source_code:
        for item in candidates:
            if (
                normalize_prefill_text(
                    item[
                        "category_code"
                    ]
                )
                ==
                source_code
            ):
                return item

    if source_name:
        for item in candidates:
            if (
                normalize_prefill_text(
                    item[
                        "category_name"
                    ]
                )
                .casefold()
                ==
                source_name
            ):
                return item

    return None


def find_prefill_subcategory(
    subcategories,
    scheme_id,
    category_id,
    source_value,
):
    if (
        prefill_status(
            source_value
        )
        !=
        STATUS_VALUE
    ):
        return None

    source_code = (
        extract_prefill_code(
            source_value
        )
    )

    source_name = (
        remove_prefill_code(
            source_value
        )
        .casefold()
    )

    candidates = []

    for item in subcategories:
        if (
            item[
                "scheme_id"
            ]
            !=
            scheme_id
        ):
            continue

        if category_id is None:
            if (
                item[
                    "category_id"
                ]
                is not None
            ):
                continue

        else:
            if (
                item[
                    "category_id"
                ]
                !=
                category_id
            ):
                continue

        candidates.append(
            item
        )

    if source_code:
        for item in candidates:
            if (
                normalize_prefill_text(
                    item[
                        "subcategory_code"
                    ]
                )
                ==
                source_code
            ):
                return item

    if source_name:
        for item in candidates:
            if (
                normalize_prefill_text(
                    item[
                        "subcategory_name"
                    ]
                )
                .casefold()
                ==
                source_name
            ):
                return item

    return None


def find_prefill_product_service(
    product_services,
    scheme_id,
    category_id,
    subcategory_id,
    source_value,
):
    if (
        prefill_status(
            source_value
        )
        !=
        STATUS_VALUE
    ):
        return None

    source_code = (
        extract_prefill_code(
            source_value
        )
    )

    source_name = (
        remove_prefill_code(
            source_value
        )
        .casefold()
    )

    candidates = []

    for item in product_services:
        if (
            item[
                "scheme_id"
            ]
            !=
            scheme_id
        ):
            continue

        if category_id is None:
            if (
                item[
                    "category_id"
                ]
                is not None
            ):
                continue

        else:
            if (
                item[
                    "category_id"
                ]
                !=
                category_id
            ):
                continue

        if subcategory_id is None:
            if (
                item[
                    "subcategory_id"
                ]
                is not None
            ):
                continue

        else:
            if (
                item[
                    "subcategory_id"
                ]
                !=
                subcategory_id
            ):
                continue

        candidates.append(
            item
        )

    if source_code:
        for item in candidates:
            if (
                normalize_prefill_text(
                    item[
                        "product_service_code"
                    ]
                )
                ==
                source_code
            ):
                return item

    if source_name:
        for item in candidates:
            if (
                normalize_prefill_text(
                    item[
                        "product_service_name"
                    ]
                )
                .casefold()
                ==
                source_name
            ):
                return item

    return None


def get_scope_review_prefill(
    conn,
    batch_id,
    issue_id,
    options,
):
    issue = (
        conn.execute(
            """
            SELECT
                issue.issue_id,
                issue.batch_id,
                issue.row_no,
                issue.field_name,
                issue.reason,

                staging.staging_id,
                staging.raw_no_lpk,
                staging.raw_scheme,
                staging.raw_category,
                staging.raw_subcategory,
                staging.raw_product_service

            FROM scope_review_issue AS issue

            JOIN scope_staging AS staging
                ON staging.staging_id
                =
                issue.staging_id

            WHERE issue.issue_id = ?
            AND issue.batch_id = ?
            """,
            (
                issue_id,
                batch_id,
            ),
        )
        .fetchone()
    )

    if issue is None:
        return None

    scheme_item = (
        find_prefill_scheme(
            options[
                "schemes"
            ],
            issue[
                "raw_scheme"
            ],
        )
    )

    scheme_id = (
        scheme_item[
            "scheme_id"
        ]
        if scheme_item
        else None
    )

    category_status = (
        prefill_status(
            issue[
                "raw_category"
            ]
        )
    )

    category_item = None

    if scheme_id is not None:
        category_item = (
            find_prefill_category(
                options[
                    "categories"
                ],
                scheme_id,
                issue[
                    "raw_category"
                ],
            )
        )

    category_id = (
        category_item[
            "category_id"
        ]
        if category_item
        else None
    )

    subcategory_status = (
        prefill_status(
            issue[
                "raw_subcategory"
            ]
        )
    )

    subcategory_item = None

    if scheme_id is not None:
        subcategory_item = (
            find_prefill_subcategory(
                options[
                    "subcategories"
                ],
                scheme_id,
                category_id,
                issue[
                    "raw_subcategory"
                ],
            )
        )

    subcategory_id = (
        subcategory_item[
            "subcategory_id"
        ]
        if subcategory_item
        else None
    )

    product_status = (
        prefill_status(
            issue[
                "raw_product_service"
            ]
        )
    )

    product_item = None

    if scheme_id is not None:
        product_item = (
            find_prefill_product_service(
                options[
                    "product_services"
                ],
                scheme_id,
                category_id,
                subcategory_id,
                issue[
                    "raw_product_service"
                ],
            )
        )

    return {
        "active":
            True,

        "batch_id":
            batch_id,

        "issue_id":
            issue_id,

        "row_no":
            issue[
                "row_no"
            ],

        "field_name":
            issue[
                "field_name"
            ],

        "reason":
            issue[
                "reason"
            ],

        "raw_no_lpk":
            issue[
                "raw_no_lpk"
            ],

        "raw_scheme":
            issue[
                "raw_scheme"
            ],

        "raw_category":
            issue[
                "raw_category"
            ],

        "raw_subcategory":
            issue[
                "raw_subcategory"
            ],

        "raw_product_service":
            issue[
                "raw_product_service"
            ],

        "scheme_id":
            (
                scheme_item[
                    "scheme_id"
                ]
                if scheme_item
                else None
            ),

        "scheme_name":
            (
                scheme_item[
                    "scheme_name"
                ]
                if scheme_item
                else remove_prefill_code(
                    issue[
                        "raw_scheme"
                    ]
                )
            ),

        "scheme_code":
            (
                scheme_item[
                    "scheme_code"
                ]
                if scheme_item
                else ""
            ),

        "category_status":
            category_status,

        "category_id":
            (
                category_item[
                    "category_id"
                ]
                if category_item
                else None
            ),

        "category_name":
            (
                category_item[
                    "category_name"
                ]
                if category_item
                else remove_prefill_code(
                    issue[
                        "raw_category"
                    ]
                )
            ),

        "category_code":
            (
                category_item[
                    "category_code"
                ]
                if category_item
                else extract_prefill_code(
                    issue[
                        "raw_category"
                    ]
                )
            ),

        "subcategory_status":
            subcategory_status,

        "subcategory_id":
            (
                subcategory_item[
                    "subcategory_id"
                ]
                if subcategory_item
                else None
            ),

        "subcategory_name":
            (
                subcategory_item[
                    "subcategory_name"
                ]
                if subcategory_item
                else remove_prefill_code(
                    issue[
                        "raw_subcategory"
                    ]
                )
            ),

        "subcategory_code":
            (
                subcategory_item[
                    "subcategory_code"
                ]
                if subcategory_item
                else extract_prefill_code(
                    issue[
                        "raw_subcategory"
                    ]
                )
            ),

        "product_service_status":
            product_status,

        "product_service_id":
            (
                product_item[
                    "product_service_id"
                ]
                if product_item
                else None
            ),

        "product_service_name":
            (
                product_item[
                    "product_service_name"
                ]
                if product_item
                else remove_prefill_code(
                    issue[
                        "raw_product_service"
                    ]
                )
            ),

        "product_service_code":
            (
                product_item[
                    "product_service_code"
                ]
                if product_item
                else extract_prefill_code(
                    issue[
                        "raw_product_service"
                    ]
                )
            ),

        "product_exists":
            (
                product_item
                is not None
            ),
    }


def resolve_scope_review_after_product_create(
    conn,
    batch_id,
    issue_id,
    path_id,
):
    issue = (
        conn.execute(
            """
            SELECT
                issue.issue_id,
                issue.staging_id

            FROM scope_review_issue AS issue

            WHERE issue.issue_id = ?
            AND issue.batch_id = ?
            AND issue.decision IS NULL
            """,
            (
                issue_id,
                batch_id,
            ),
        )
        .fetchone()
    )

    if issue is None:
        return

    staging_id = (
        issue[
            "staging_id"
        ]
    )

    conn.execute(
        """
        UPDATE scope_staging

        SET
            product_path_id = ?,
            updated_at = CURRENT_TIMESTAMP

        WHERE staging_id = ?
        """,
        (
            path_id,
            staging_id,
        ),
    )

    conn.execute(
        """
        UPDATE scope_review_issue

        SET
            decision = 'RESOLVED',
            resolved_value = ?,
            decision_note =
                'Master Produk ditambahkan dari Review Data dan issue diselesaikan otomatis.',
            resolved_at = CURRENT_TIMESTAMP,
            resolved_by = 'Admin',
            updated_at = CURRENT_TIMESTAMP

        WHERE issue_id = ?
        """,
        (
            f"product_path_id={path_id}",
            issue_id,
        ),
    )

    unresolved_issue = (
        conn.execute(
            """
            SELECT COUNT(*) AS total

            FROM scope_review_issue

            WHERE staging_id = ?
            AND decision IS NULL
            """,
            (
                staging_id,
            ),
        )
        .fetchone()[
            "total"
        ]
    )

    unresolved_blocking = (
        conn.execute(
            """
            SELECT COUNT(*) AS total

            FROM scope_review_issue

            WHERE staging_id = ?
            AND is_blocking = 1
            AND decision IS NULL
            """,
            (
                staging_id,
            ),
        )
        .fetchone()[
            "total"
        ]
    )

    pending_correction = (
        conn.execute(
            """
            SELECT COUNT(*) AS total

            FROM scope_correction_log

            WHERE staging_id = ?
            AND approval_status = 'PENDING'
            """,
            (
                staging_id,
            ),
        )
        .fetchone()[
            "total"
        ]
    )

    if unresolved_issue > 0:
        row_status = "Need Review"

    elif pending_correction > 0:
        row_status = "Corrected"

    else:
        row_status = "Valid"

    conn.execute(
        """
        UPDATE scope_staging

        SET
            row_status = ?,
            blocking_issue_count = ?,
            updated_at = CURRENT_TIMESTAMP

        WHERE staging_id = ?
        """,
        (
            row_status,
            unresolved_blocking,
            staging_id,
        ),
    )

    recalculate_scope_batch(
        conn,
        batch_id,
    )


# ============================================================
# CREATE PAGE
# ============================================================

@router.get(
    "/products/create",
    response_class=HTMLResponse,
)
def product_create_page(
    request: Request,
    status: str = "",
    message: str = "",
    source: str = "",
    batch_id: int | None = None,
    issue_id: int | None = None,
):

    options = (
        get_product_form_options()
    )

    review_prefill = {
        "active":
            False,
    }

    if (
        source
        ==
        "scope_review"
        and
        batch_id is not None
        and
        issue_id is not None
    ):
        conn = get_connection()

        try:
            prefill = (
                get_scope_review_prefill(
                    conn,
                    batch_id,
                    issue_id,
                    options,
                )
            )

            if prefill:
                review_prefill = (
                    prefill
                )

        finally:
            conn.close()

    return templates.TemplateResponse(
        request=request,

        name="products/create.html",

        context={
            "status":
                status,

            "message":
                message,

            "schemes":
                options[
                    "schemes"
                ],

            "categories":
                options[
                    "categories"
                ],

            "subcategories":
                options[
                    "subcategories"
                ],

            "product_services":
                options[
                    "product_services"
                ],

            "review_prefill":
                review_prefill,
        },
    )


# ============================================================
# CREATE PROCESS
# ============================================================

@router.post(
    "/products/create"
)
def product_create_process(

    scheme_name: str = Form(...),
    scheme_code: str = Form(...),

    category_status: str = Form(...),
    category_name: str = Form(""),
    category_code: str = Form(""),

    subcategory_status: str = Form(...),
    subcategory_name: str = Form(""),
    subcategory_code: str = Form(""),

    product_service_status: str = Form(...),
    product_service_name: str = Form(""),
    product_service_code: str = Form(""),

    source: str = Form(""),
    batch_id: int | None = Form(None),
    issue_id: int | None = Form(None),
):

    conn = get_connection()


    try:

        validate_product_form(
            scheme_name,
            scheme_code,

            category_status,
            category_name,
            category_code,

            subcategory_status,
            subcategory_name,
            subcategory_code,

            product_service_status,
            product_service_name,
            product_service_code,
        )


        conn.execute(
            "BEGIN"
        )


        # ====================================================
        # SCHEME
        # ====================================================

        scheme_id = (
            get_or_create_scheme(
                conn,
                scheme_code,
                scheme_name,
            )
        )


        # ====================================================
        # CATEGORY
        # ====================================================

        category_id = None


        if category_status == STATUS_VALUE:

            category_id = (
                get_or_create_category(
                    conn,
                    scheme_id,
                    category_code,
                    category_name,
                )
            )


        # ====================================================
        # SUBCATEGORY
        # ====================================================

        subcategory_id = None


        if subcategory_status == STATUS_VALUE:

            subcategory_id = (
                get_or_create_subcategory(
                    conn,
                    scheme_id,
                    category_id,
                    subcategory_code,
                    subcategory_name,
                )
            )


        # ====================================================
        # DUPLICATE ID JENIS
        # ====================================================

        if product_service_status == STATUS_VALUE:

            duplicate = (
                find_duplicate_product_code(
                    conn=conn,

                    scheme_id=scheme_id,

                    category_id=category_id,

                    subcategory_id=(
                        subcategory_id
                    ),

                    product_service_code=(
                        product_service_code
                    ),
                )
            )


            if duplicate:

                raise ValueError(
                    (
                        f"ID Jenis "
                        f"'{product_service_code}' "
                        "sudah digunakan pada "
                        "Skema/Kategori/Sub Kategori "
                        "yang sama oleh Jenis "
                        f"'{duplicate['product_service_name']}'."
                    )
                )


        # ====================================================
        # PRODUCT SERVICE
        # ====================================================

        product_service_id = None


        if product_service_status == STATUS_VALUE:

            product_service_id = (
                get_or_create_product_service(
                    conn,
                    scheme_id,
                    category_id,
                    subcategory_id,
                    product_service_code,
                    product_service_name,
                )
            )


        # ====================================================
        # CHECK EXACT PATH
        # ====================================================

        query = """
            SELECT path_id

            FROM master_product_path

            WHERE scheme_id = ?

            AND category_status = ?

            AND subcategory_status = ?

            AND product_service_status = ?
        """


        params = [
            scheme_id,
            category_status,
            subcategory_status,
            product_service_status,
        ]


        if category_id is None:

            query += """
                AND category_id IS NULL
            """

        else:

            query += """
                AND category_id = ?
            """

            params.append(
                category_id
            )


        if subcategory_id is None:

            query += """
                AND subcategory_id IS NULL
            """

        else:

            query += """
                AND subcategory_id = ?
            """

            params.append(
                subcategory_id
            )


        if product_service_id is None:

            query += """
                AND product_service_id IS NULL
            """

        else:

            query += """
                AND product_service_id = ?
            """

            params.append(
                product_service_id
            )


        query += """
            LIMIT 1
        """


        duplicate_path = (
            conn.execute(
                query,
                params,
            )
            .fetchone()
        )


        if duplicate_path:

            raise ValueError(
                (
                    "Data Master Produk "
                    "dengan jalur yang sama "
                    "sudah tersedia."
                )
            )


        # ====================================================
        # INSERT
        # ====================================================

        path_cursor = conn.execute(
            """
            INSERT INTO master_product_path (
                scheme_id,
                category_id,
                subcategory_id,
                product_service_id,

                category_status,
                subcategory_status,
                product_service_status,

                created_at,
                updated_at
            )

            VALUES (
                ?,
                ?,
                ?,
                ?,

                ?,
                ?,
                ?,

                CURRENT_TIMESTAMP,
                CURRENT_TIMESTAMP
            )
            """,
            (
                scheme_id,
                category_id,
                subcategory_id,
                product_service_id,

                category_status,
                subcategory_status,
                product_service_status,
            ),
        )


        path_id = (
            path_cursor.lastrowid
        )


        if (
            source
            ==
            "scope_review"
            and
            batch_id is not None
            and
            issue_id is not None
        ):
            resolve_scope_review_after_product_create(
                conn=conn,
                batch_id=batch_id,
                issue_id=issue_id,
                path_id=path_id,
            )


        conn.commit()


    except Exception as exc:

        conn.rollback()


        query_params = {
            "status":
                "failed",

            "message":
                str(exc),
        }

        if (
            source
            ==
            "scope_review"
        ):
            query_params[
                "source"
            ] = "scope_review"

            if batch_id is not None:
                query_params[
                    "batch_id"
                ] = batch_id

            if issue_id is not None:
                query_params[
                    "issue_id"
                ] = issue_id

        query = urlencode(
            query_params
        )


        return RedirectResponse(
            url=(
                "/products/create?"
                + query
            ),

            status_code=303,
        )


    finally:

        conn.close()


    if (
        source
        ==
        "scope_review"
        and
        batch_id is not None
    ):
        query = urlencode(
            {
                "result":
                    "success",

                "message":
                    (
                        "Master Produk berhasil ditambahkan "
                        "dan masalah pada Review Data "
                        "berhasil diselesaikan otomatis."
                    ),
            }
        )

        return RedirectResponse(
            url=(
                f"/scopes/batches/"
                f"{batch_id}/review?"
                + query
            ),

            status_code=303,
        )


    query = urlencode(
        {
            "result":
                "success",

            "message":
                "Master Produk berhasil ditambahkan.",
        }
    )


    return RedirectResponse(
        url=(
            "/products?"
            + query
        ),

        status_code=303,
    )


# ============================================================
# EDIT PAGE
# ============================================================

@router.get(
    "/products/{path_id}/edit",
    response_class=HTMLResponse,
)
def product_edit_page(
    request: Request,
    path_id: int,
    status: str = "",
    message: str = "",
):

    conn = get_connection()


    try:

        product = (
            conn.execute(
                """
                SELECT
                    path.*,

                    scheme.scheme_code,
                    scheme.scheme_name,

                    category.category_code,
                    category.category_name,

                    subcategory.subcategory_code,
                    subcategory.subcategory_name,

                    product.product_service_code,
                    product.product_service_name

                FROM master_product_path AS path

                JOIN master_scheme AS scheme
                    ON scheme.scheme_id
                    =
                    path.scheme_id

                LEFT JOIN master_category AS category
                    ON category.category_id
                    =
                    path.category_id

                LEFT JOIN master_subcategory AS subcategory
                    ON subcategory.subcategory_id
                    =
                    path.subcategory_id

                LEFT JOIN master_product_service AS product
                    ON product.product_service_id
                    =
                    path.product_service_id

                WHERE path.path_id = ?
                """,
                (
                    path_id,
                ),
            )
            .fetchone()
        )


        if product is None:

            raise HTTPException(
                status_code=404,

                detail=(
                    "Data Master Produk "
                    "tidak ditemukan."
                ),
            )


    finally:

        conn.close()


    return templates.TemplateResponse(
        request=request,

        name="products/edit.html",

        context={
            "product":
                product,

            "status":
                status,

            "message":
                message,
        },
    )


# ============================================================
# EDIT PROCESS
# ============================================================

@router.post(
    "/products/{path_id}/edit"
)
def product_edit_process(

    path_id: int,

    scheme_name: str = Form(...),
    scheme_code: str = Form(...),

    category_status: str = Form(...),
    category_name: str = Form(""),
    category_code: str = Form(""),

    subcategory_status: str = Form(...),
    subcategory_name: str = Form(""),
    subcategory_code: str = Form(""),

    product_service_status: str = Form(...),
    product_service_name: str = Form(""),
    product_service_code: str = Form(""),
):

    conn = get_connection()


    try:

        existing_path = (
            conn.execute(
                """
                SELECT path_id

                FROM master_product_path

                WHERE path_id = ?
                """,
                (
                    path_id,
                ),
            )
            .fetchone()
        )


        if existing_path is None:

            raise HTTPException(
                status_code=404,

                detail=(
                    "Data Master Produk "
                    "tidak ditemukan."
                ),
            )


        validate_product_form(
            scheme_name,
            scheme_code,

            category_status,
            category_name,
            category_code,

            subcategory_status,
            subcategory_name,
            subcategory_code,

            product_service_status,
            product_service_name,
            product_service_code,
        )


        conn.execute(
            "BEGIN"
        )


        scheme_id = (
            get_or_create_scheme(
                conn,
                scheme_code,
                scheme_name,
            )
        )


        category_id = None


        if category_status == STATUS_VALUE:

            category_id = (
                get_or_create_category(
                    conn,
                    scheme_id,
                    category_code,
                    category_name,
                )
            )


        subcategory_id = None


        if subcategory_status == STATUS_VALUE:

            subcategory_id = (
                get_or_create_subcategory(
                    conn,
                    scheme_id,
                    category_id,
                    subcategory_code,
                    subcategory_name,
                )
            )


        if product_service_status == STATUS_VALUE:

            duplicate = (
                find_duplicate_product_code(
                    conn=conn,

                    scheme_id=scheme_id,

                    category_id=category_id,

                    subcategory_id=(
                        subcategory_id
                    ),

                    product_service_code=(
                        product_service_code
                    ),

                    exclude_path_id=(
                        path_id
                    ),
                )
            )


            if duplicate:

                raise ValueError(
                    (
                        f"ID Jenis "
                        f"'{product_service_code}' "
                        "sudah digunakan pada "
                        "Skema/Kategori/Sub Kategori "
                        "yang sama oleh Jenis "
                        f"'{duplicate['product_service_name']}'."
                    )
                )


        product_service_id = None


        if product_service_status == STATUS_VALUE:

            product_service_id = (
                get_or_create_product_service(
                    conn,
                    scheme_id,
                    category_id,
                    subcategory_id,
                    product_service_code,
                    product_service_name,
                )
            )


        conn.execute(
            """
            UPDATE master_product_path

            SET
                scheme_id = ?,
                category_id = ?,
                subcategory_id = ?,
                product_service_id = ?,

                category_status = ?,
                subcategory_status = ?,
                product_service_status = ?,

                updated_at = CURRENT_TIMESTAMP

            WHERE path_id = ?
            """,
            (
                scheme_id,
                category_id,
                subcategory_id,
                product_service_id,

                category_status,
                subcategory_status,
                product_service_status,

                path_id,
            ),
        )


        conn.commit()


    except HTTPException:

        conn.rollback()

        raise


    except Exception as exc:

        conn.rollback()


        query = urlencode(
            {
                "status":
                    "failed",

                "message":
                    str(exc),
            }
        )


        return RedirectResponse(
            url=(
                f"/products/{path_id}/edit?"
                + query
            ),

            status_code=303,
        )


    finally:

        conn.close()


    query = urlencode(
        {
            "result":
                "success",

            "message":
                "Master Produk berhasil diperbarui.",
        }
    )


    return RedirectResponse(
        url=(
            "/products?"
            + query
        ),

        status_code=303,
    )


# ============================================================
# DELETE
# ============================================================

@router.post(
    "/products/{path_id}/delete"
)
def product_delete(
    path_id: int,
):

    conn = get_connection()


    try:

        existing = (
            conn.execute(
                """
                SELECT path_id

                FROM master_product_path

                WHERE path_id = ?
                """,
                (
                    path_id,
                ),
            )
            .fetchone()
        )


        if existing is None:

            query = urlencode(
                {
                    "result":
                        "failed",

                    "message":
                        "Data Master Produk tidak ditemukan.",
                }
            )


            return RedirectResponse(
                url=(
                    "/products?"
                    + query
                ),

                status_code=303,
            )


        conn.execute(
            """
            DELETE FROM master_product_path

            WHERE path_id = ?
            """,
            (
                path_id,
            ),
        )


        conn.commit()


    except Exception as exc:

        conn.rollback()


        query = urlencode(
            {
                "result":
                    "failed",

                "message":
                    str(exc),
            }
        )


        return RedirectResponse(
            url=(
                "/products?"
                + query
            ),

            status_code=303,
        )


    finally:

        conn.close()


    query = urlencode(
        {
            "result":
                "success",

            "message":
                "Master Produk berhasil dihapus.",
        }
    )


    return RedirectResponse(
        url=(
            "/products?"
            + query
        ),

        status_code=303,
    )


# ============================================================
# IMPORT PAGE
# ============================================================

@router.get(
    "/products/import",
    response_class=HTMLResponse,
)
def product_import_page(
    request: Request,
    status: str = "",
    message: str = "",
    import_mode: str = "",
    total_rows: int = 0,
    inserted_path_count: int = 0,
    existing_path_count: int = 0,
):

    import_result = None
    success_message = None
    error_message = None


    if status == "success":

        success_message = (
            message
            or
            "Master Produk berhasil di-import."
        )


        import_result = {
            "import_mode":
                import_mode,

            "total_rows":
                total_rows,

            "inserted_path_count":
                inserted_path_count,

            "existing_path_count":
                existing_path_count,
        }


    elif status == "failed":

        error_message = (
            message
            or
            "Import Master Produk gagal."
        )


    return templates.TemplateResponse(
        request=request,

        name="products/import.html",

        context={
            "success_message":
                success_message,

            "error_message":
                error_message,

            "import_result":
                import_result,
        },
    )


# ============================================================
# IMPORT PROCESS
# ============================================================

@router.post(
    "/products/import"
)
async def product_import_process(

    file: UploadFile = File(...),

    import_mode: str = Form(...),
):

    try:

        if not file.filename:

            raise ValueError(
                "Nama file tidak ditemukan."
            )


        allowed_extensions = (
            ".xlsx",
            ".xls",
            ".csv",
        )


        if not file.filename.lower().endswith(
            allowed_extensions
        ):

            raise ValueError(
                "Format file tidak didukung. "
                "Gunakan .xlsx, .xls, atau .csv."
            )


        file_content = (
            await file.read()
        )


        if not file_content:

            raise ValueError(
                "File yang di-upload kosong."
            )


        result = import_products(
            file_content=file_content,
            filename=file.filename,
            import_mode=import_mode,
        )


        mode_name = (
            "Upload Ulang Semua Data"

            if import_mode == "replace"

            else "Tambah Data"
        )


        query = urlencode(
            {
                "status":
                    "success",

                "message":
                    (
                        "Import Master Produk berhasil "
                        f"dengan mode {mode_name}."
                    ),

                "import_mode":
                    result["import_mode"],

                "total_rows":
                    result["total_rows"],

                "inserted_path_count":
                    result[
                        "inserted_path_count"
                    ],

                "existing_path_count":
                    result[
                        "existing_path_count"
                    ],
            }
        )


        return RedirectResponse(
            url=(
                "/products/import?"
                + query
            ),

            status_code=303,
        )


    except Exception as exc:

        query = urlencode(
            {
                "status":
                    "failed",

                "message":
                    str(exc),
            }
        )


        return RedirectResponse(
            url=(
                "/products/import?"
                + query
            ),

            status_code=303,
        )


# ============================================================
# DOWNLOAD TEMPLATE
# ============================================================

@router.get(
    "/products/template"
)
def download_product_template():

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = (
        "Master Data Produk (Clean)"
    )


    for (
        column_index,
        column_name,
    ) in enumerate(
        PRODUCT_COLUMNS,
        start=1,
    ):

        cell = worksheet.cell(
            row=1,
            column=column_index,
            value=column_name,
        )

        cell.font = Font(
            bold=True
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )


    worksheet.freeze_panes = (
        "A2"
    )


    guide_sheet = (
        workbook.create_sheet(
            "Panduan"
        )
    )


    guide_sheet.append(
        [
            "Aturan",
            "Makna",
        ]
    )


    guide_sheet.append(
        [
            "Ada nilai",
            "VALUE - data tersedia",
        ]
    )


    guide_sheet.append(
        [
            "-",
            (
                "NOT_APPLICABLE - "
                "level tidak berlaku"
            ),
        ]
    )


    guide_sheet.append(
        [
            "Kosong",
            (
                "MISSING - level ada "
                "tetapi belum diisi"
            ),
        ]
    )


    output = BytesIO()

    workbook.save(
        output
    )

    output.seek(0)


    return StreamingResponse(
        output,

        media_type=(
            "application/"
            "vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),

        headers={
            "Content-Disposition":
                (
                    "attachment; "
                    "filename="
                    "Template_Master_Produk.xlsx"
                )
        },
    )


# ============================================================
# EXPORT EXCEL
# ============================================================

@router.get(
    "/products/export/xlsx"
)
def export_products_xlsx():

    rows = (
        get_product_export_data()
    )


    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = (
        "Master Data Produk (Clean)"
    )


    for (
        column_index,
        column_name,
    ) in enumerate(
        PRODUCT_COLUMNS,
        start=1,
    ):

        cell = worksheet.cell(
            row=1,
            column=column_index,
            value=column_name,
        )

        cell.font = Font(
            bold=True
        )


    for (
        row_index,
        item,
    ) in enumerate(
        rows,
        start=2,
    ):

        for (
            column_index,
            column_name,
        ) in enumerate(
            PRODUCT_COLUMNS,
            start=1,
        ):

            worksheet.cell(
                row=row_index,
                column=column_index,
                value=item[
                    column_name
                ],
            )


    worksheet.freeze_panes = (
        "A2"
    )


    worksheet.auto_filter.ref = (
        worksheet.dimensions
    )


    output = BytesIO()

    workbook.save(
        output
    )

    output.seek(0)


    return StreamingResponse(
        output,

        media_type=(
            "application/"
            "vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),

        headers={
            "Content-Disposition":
                (
                    "attachment; "
                    "filename="
                    "Master_Data_Produk.xlsx"
                )
        },
    )


# ============================================================
# EXPORT CSV
# ============================================================

@router.get(
    "/products/export/csv"
)
def export_products_csv():

    rows = (
        get_product_export_data()
    )


    output = StringIO(
        newline=""
    )


    writer = csv.DictWriter(
        output,

        fieldnames=(
            PRODUCT_COLUMNS
        ),
    )


    writer.writeheader()


    for item in rows:

        writer.writerow(
            item
        )


    csv_content = (
        "\ufeff"
        + output.getvalue()
    )


    csv_bytes = (
        csv_content.encode(
            "utf-8"
        )
    )


    return StreamingResponse(
        iter(
            [
                csv_bytes
            ]
        ),

        media_type=(
            "text/csv; charset=utf-8"
        ),

        headers={
            "Content-Disposition":
                (
                    "attachment; "
                    "filename="
                    "Master_Data_Produk.csv"
                )
        },
    )