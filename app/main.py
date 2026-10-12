from pathlib import Path



from urllib.parse import urlencode



from io import BytesIO, StringIO



from app.routes.products import (



    router as products_router,



)



from app.routes.certification_schemes import (



    router as certification_schemes_router,



)




from app.routes.standards import (

    router as standards_router,

)

from app.routes.scopes import (

    router as scopes_router,

)


import csv





from fastapi import (



    FastAPI,



    Request,



    HTTPException,



    UploadFile,



    File,



    Form,



)





from fastapi.responses import (



    HTMLResponse,



    RedirectResponse,



    StreamingResponse,



)





from fastapi.templating import (



    Jinja2Templates,



)





from fastapi.staticfiles import (



    StaticFiles,



)





from openpyxl import Workbook





from openpyxl.styles import (



    Font,



    Alignment,



)





from app.database import (



    get_connection,



)





from app.services.institution_import import (



    import_institutions,



)





# ============================================================



# APP



# ============================================================





app = FastAPI(



    title=(



        "Sistem Database "



        "Lembaga Sertifikasi"



    )



)





app.include_router(



    certification_schemes_router



)





app.include_router(



    products_router



)






app.include_router(

    standards_router

)

app.include_router(

    scopes_router

)


# ============================================================



# PATH



# ============================================================





BASE_DIR = (



    Path(__file__)



    .resolve()



    .parent



)





TEMPLATE_DIR = (



    BASE_DIR



    / "templates"



)





STATIC_DIR = (



    BASE_DIR



    / "static"



)





# ============================================================



# TEMPLATE



# ============================================================





templates = Jinja2Templates(



    directory=str(



        TEMPLATE_DIR



    )



)





# ============================================================



# STATIC



# ============================================================





app.mount(



    "/static",



    StaticFiles(



        directory=str(



            STATIC_DIR



        )



    ),



    name="static",



)





# ============================================================



# FORMAT MASTER LEMBAGA



# ============================================================





INSTITUTION_TEMPLATE_COLUMNS = [



    "ID_LSPR",



    "Name_LSPR",



    "Date_Start_Accreditation",



    "Date_End_Accreditation",



    "Status_Accreditation",



    "Amandemen",



    "Address",



    "Kab_Kota",



    "Propinsi",



    "Kode_Pos",



    "Telp",



    "Fax",



    "Email",



    "CP",



    "NoHpCP",



]





# ============================================================



# HELPER AMANDEMEN EXPORT



# ============================================================





def build_amendment_export_text(amendments):





    if not amendments:



        return ""





    results = []





    for amendment in amendments:





        amendment_no = amendment["amendment_no"]



        amendment_date = amendment["amendment_date"]





        if amendment_date:





            text = (



                f"Amandemen ke-{amendment_no} "



                f"({amendment_date})"



            )





        else:





            text = (



                f"Amandemen ke-{amendment_no}"



            )





        results.append(text)





    return "\n".join(results)





# ============================================================



# HELPER DATA EXPORT



# ============================================================





def get_institution_export_data():





    conn = get_connection()





    try:





        institutions = (



            conn.execute(



                """



                SELECT *





                FROM institution





                WHERE source_status = 'AVAILABLE'





                ORDER BY id_lspro ASC



                """



            )



            .fetchall()



        )





        export_rows = []





        for institution in institutions:





            amendments = (



                conn.execute(



                    """



                    SELECT



                        amendment_no,



                        amendment_date





                    FROM institution_amendment





                    WHERE institution_id = ?





                    ORDER BY amendment_no ASC



                    """,



                    (



                        institution[



                            "institution_id"



                        ],



                    ),



                )



                .fetchall()



            )





            amendment_text = (



                build_amendment_export_text(



                    amendments



                )



            )





            export_rows.append(



                {



                    "ID_LSPR":



                        institution[



                            "id_lspro"



                        ],





                    "Name_LSPR":



                        institution[



                            "name_lspro"



                        ],





                    "Date_Start_Accreditation":



                        institution[



                            "date_start_accreditation"



                        ],





                    "Date_End_Accreditation":



                        institution[



                            "date_end_accreditation"



                        ],





                    "Status_Accreditation":



                        institution[



                            "status_accreditation"



                        ],





                    "Amandemen":



                        amendment_text,





                    "Address":



                        institution[



                            "address"



                        ],





                    "Kab_Kota":



                        institution[



                            "kab_kota"



                        ],





                    "Propinsi":



                        institution[



                            "propinsi"



                        ],





                    "Kode_Pos":



                        institution[



                            "kode_pos"



                        ],





                    "Telp":



                        institution[



                            "telp"



                        ],





                    "Fax":



                        institution[



                            "fax"



                        ],





                    "Email":



                        institution[



                            "email"



                        ],





                    "CP":



                        institution[



                            "cp"



                        ],





                    "NoHpCP":



                        institution[



                            "no_hp_cp"



                        ],



                }



            )





        return export_rows





    finally:





        conn.close()





# ============================================================

# DASHBOARD

# ============================================================



@app.get(

    "/",

    response_class=HTMLResponse,

)

def dashboard(

    request: Request,

):

    conn = get_connection()



    try:

        # ----------------------------------------------------

        # RINGKASAN MASTER LEMBAGA

        # ----------------------------------------------------



        total_institution = (

            conn.execute(

                """

                SELECT COUNT(*) AS total

                FROM institution

                WHERE source_status = 'AVAILABLE'

                """

            )

            .fetchone()["total"]

        )



        active_count = (

            conn.execute(

                """

                SELECT COUNT(*) AS total

                FROM institution

                WHERE source_status = 'AVAILABLE'

                AND LOWER(COALESCE(status_accreditation, '')) = 'active'

                """

            )

            .fetchone()["total"]

        )



        withdrawn_count = (

            conn.execute(

                """

                SELECT COUNT(*) AS total

                FROM institution

                WHERE source_status = 'AVAILABLE'

                AND LOWER(COALESCE(status_accreditation, '')) = 'withdrawn'

                """

            )

            .fetchone()["total"]

        )



        other_status_count = max(

            total_institution

            - active_count

            - withdrawn_count,

            0,

        )



        total_amendment = (

            conn.execute(

                """

                SELECT COUNT(*) AS total

                FROM institution_amendment

                """

            )

            .fetchone()["total"]

        )



        # ----------------------------------------------------

        # RINGKASAN MASTER PRODUK & SKEMA SERTIFIKASI

        # ----------------------------------------------------



        total_product_paths = (

            conn.execute(

                """

                SELECT COUNT(*) AS total

                FROM master_product_path

                """

            )

            .fetchone()["total"]

        )



        total_product_schemes = (

            conn.execute(

                """

                SELECT COUNT(*) AS total

                FROM master_scheme

                """

            )

            .fetchone()["total"]

        )



        total_certification_schemes = (

            conn.execute(

                """

                SELECT COUNT(*) AS total

                FROM master_certification_scheme

                """

            )

            .fetchone()["total"]

        )



        # ----------------------------------------------------

        # DISTRIBUSI LEMBAGA PER PROVINSI

        # ----------------------------------------------------



        province_rows = (

            conn.execute(

                """

                SELECT

                    CASE

                        WHEN TRIM(COALESCE(propinsi, '')) = ''

                            THEN 'Belum Diisi'

                        ELSE TRIM(propinsi)

                    END AS province,

                    COUNT(*) AS total



                FROM institution



                WHERE source_status = 'AVAILABLE'



                GROUP BY

                    CASE

                        WHEN TRIM(COALESCE(propinsi, '')) = ''

                            THEN 'Belum Diisi'

                        ELSE TRIM(propinsi)

                    END



                ORDER BY

                    total DESC,

                    province ASC



                LIMIT 8

                """

            )

            .fetchall()

        )



        province_data = [

            {

                "province": row["province"],

                "total": row["total"],

            }

            for row in province_rows

        ]



        max_province_total = (

            max(

                (

                    item["total"]

                    for item in province_data

                ),

                default=1,

            )

        )



        for item in province_data:

            item["percentage"] = round(

                (

                    item["total"]

                    / max_province_total

                )

                * 100,

                2,

            )



        # ----------------------------------------------------

        # STATUS LEMBAGA

        # ----------------------------------------------------



        status_rows = (

            conn.execute(

                """

                SELECT

                    CASE

                        WHEN TRIM(

                            COALESCE(

                                status_accreditation,

                                ''

                            )

                        ) = ''

                            THEN 'Belum Diisi'

                        ELSE TRIM(status_accreditation)

                    END AS status_name,

                    COUNT(*) AS total



                FROM institution



                WHERE source_status = 'AVAILABLE'



                GROUP BY

                    CASE

                        WHEN TRIM(

                            COALESCE(

                                status_accreditation,

                                ''

                            )

                        ) = ''

                            THEN 'Belum Diisi'

                        ELSE TRIM(status_accreditation)

                    END



                ORDER BY

                    total DESC,

                    status_name ASC

                """

            )

            .fetchall()

        )



        status_data = [

            {

                "status_name": row["status_name"],

                "total": row["total"],

                "percentage": (

                    round(

                        (

                            row["total"]

                            / total_institution

                        )

                        * 100,

                        1,

                    )

                    if total_institution

                    else 0

                ),

            }

            for row in status_rows

        ]



        # ----------------------------------------------------

        # DATA LEMBAGA TERBARU

        # ----------------------------------------------------



        recent_institutions = (

            conn.execute(

                """

                SELECT

                    institution_id,

                    id_lspro,

                    name_lspro,

                    status_accreditation,

                    propinsi



                FROM institution



                WHERE source_status = 'AVAILABLE'



                ORDER BY institution_id DESC



                LIMIT 5

                """

            )

            .fetchall()

        )



    finally:

        conn.close()



    return templates.TemplateResponse(

        request=request,

        name="dashboard.html",

        context={

            "total_institution": total_institution,

            "active_count": active_count,

            "withdrawn_count": withdrawn_count,

            "other_status_count": other_status_count,

            "total_amendment": total_amendment,

            "total_product_paths": total_product_paths,

            "total_product_schemes": total_product_schemes,

            "total_certification_schemes": total_certification_schemes,

            "province_data": province_data,

            "status_data": status_data,

            "recent_institutions": recent_institutions,

        },

    )





# ============================================================



# IMPORT PAGE



# ============================================================





@app.get(



    "/institutions/import",



    response_class=HTMLResponse,



)



def institution_import_page(



    request: Request,



    status: str = "",



    message: str = "",



    import_mode: str = "",



    total_rows: int = 0,



    inserted_count: int = 0,



    updated_count: int = 0,



    missing_count: int = 0,



):





    success_message = None



    error_message = None



    import_result = None





    if status == "success":





        success_message = (



            message



            or



            "Master Lembaga berhasil di-import."



        )





        import_result = {



            "import_mode":



                import_mode,





            "total_rows":



                total_rows,





            "inserted_count":



                inserted_count,





            "updated_count":



                updated_count,





            "missing_count":



                missing_count,



        }





    elif status == "failed":





        error_message = (



            message



            or



            "Import Master Lembaga gagal."



        )





    return templates.TemplateResponse(



        request=request,



        name="institutions/import.html",



        context={



            "success_message":



                success_message,





            "error_message":



                error_message,





            "import_result":



                import_result,



        },



    )





# ============================================================



# IMPORT PROCESS



# ============================================================





@app.post(



    "/institutions/import"



)



async def institution_import_process(





    file: UploadFile = File(...),





    import_mode: str = Form(...),



):





    try:





        if not file.filename:





            raise ValueError(



                "Nama file tidak ditemukan."



            )





        filename_lower = (



            file.filename.lower()



        )





        allowed_extensions = (



            ".xlsx",



            ".xls",



            ".csv",



        )





        if not filename_lower.endswith(



            allowed_extensions



        ):





            raise ValueError(



                "Format file tidak didukung. "



                "Gunakan .xlsx, .xls, atau .csv."



            )





        if import_mode not in {



            "append",



            "replace",



        }:





            raise ValueError(



                "Mode import tidak valid."



            )





        file_content = (



            await file.read()



        )





        if not file_content:





            raise ValueError(



                "File yang di-upload kosong."



            )





        result = import_institutions(



            file_content=file_content,



            filename=file.filename,



            import_mode=import_mode,



        )





        if import_mode == "replace":





            mode_name = (



                "Upload Ulang Semua Data"



            )





        else:





            mode_name = (



                "Tambah Data"



            )





        query = urlencode(



            {



                "status":



                    "success",





                "message":



                    (



                        "Import berhasil menggunakan "



                        f"mode {mode_name}."



                    ),





                "import_mode":



                    result[



                        "import_mode"



                    ],





                "total_rows":



                    result[



                        "total_rows"



                    ],





                "inserted_count":



                    result[



                        "inserted_count"



                    ],





                "updated_count":



                    result[



                        "updated_count"



                    ],





                "missing_count":



                    result[



                        "missing_count"



                    ],



            }



        )





        return RedirectResponse(



            url="/institutions/import?" + query,



            status_code=303,



        )





    except Exception as exc:





        query = urlencode(



            {



                "status":



                    "failed",





                "message":



                    str(exc),



            }



        )





        return RedirectResponse(



            url="/institutions/import?" + query,



            status_code=303,



        )





# ============================================================



# DOWNLOAD TEMPLATE



# ============================================================





@app.get(



    "/institutions/template"



)



def download_institution_template():





    workbook = Workbook()





    worksheet = workbook.active





    worksheet.title = (



        "Master Lembaga"



    )





    for (



        column_index,



        column_name,



    ) in enumerate(



        INSTITUTION_TEMPLATE_COLUMNS,



        start=1,



    ):





        cell = worksheet.cell(



            row=1,



            column=column_index,



            value=column_name,



        )





        cell.font = Font(



            bold=True



        )





        cell.alignment = Alignment(



            horizontal="center",



            vertical="center",



        )





    worksheet.freeze_panes = "A2"





    widths = {



        "A": 20,



        "B": 35,



        "C": 24,



        "D": 24,



        "E": 24,



        "F": 45,



        "G": 45,



        "H": 25,



        "I": 25,



        "J": 15,



        "K": 22,



        "L": 22,



        "M": 35,



        "N": 30,



        "O": 22,



    }





    for (



        column_letter,



        width,



    ) in widths.items():





        worksheet.column_dimensions[



            column_letter



        ].width = width





    guide_sheet = (



        workbook.create_sheet(



            "Panduan"



        )



    )





    guide_sheet.append(



        [



            "Nama Kolom",



            "Isi",



            "Keterangan",



        ]



    )





    guide_data = [



        (



            "ID_LSPR",



            "Wajib",



            (



                "ID unik lembaga. "



                "Contoh: LSPr-001-IDN"



            ),



        ),





        (



            "Name_LSPR",



            "Wajib",



            "Nama lembaga sertifikasi.",



        ),





        (



            "Date_Start_Accreditation",



            "Opsional",



            "Tanggal mulai akreditasi.",



        ),





        (



            "Date_End_Accreditation",



            "Opsional",



            "Tanggal akhir akreditasi.",



        ),





        (



            "Status_Accreditation",



            "Opsional",



            "Status akreditasi.",



        ),





        (



            "Amandemen",



            "Opsional",



            (



                "Contoh: "



                "Amandemen ke-1 "



                "(20 Februari 2025)"



            ),



        ),





        (



            "Address",



            "Opsional",



            "Alamat lembaga.",



        ),





        (



            "Kab_Kota",



            "Opsional",



            "Kabupaten atau kota.",



        ),





        (



            "Propinsi",



            "Opsional",



            "Provinsi.",



        ),





        (



            "Kode_Pos",



            "Opsional",



            "Kode pos.",



        ),





        (



            "Telp",



            "Opsional",



            "Nomor telepon.",



        ),





        (



            "Fax",



            "Opsional",



            "Nomor fax.",



        ),





        (



            "Email",



            "Opsional",



            "Email lembaga.",



        ),





        (



            "CP",



            "Opsional",



            "Contact Person.",



        ),





        (



            "NoHpCP",



            "Opsional",



            "Nomor HP Contact Person.",



        ),



    ]





    for row in guide_data:





        guide_sheet.append(row)





    for cell in guide_sheet[1]:





        cell.font = Font(



            bold=True



        )





    guide_sheet.column_dimensions[



        "A"



    ].width = 30





    guide_sheet.column_dimensions[



        "B"



    ].width = 15





    guide_sheet.column_dimensions[



        "C"



    ].width = 60





    output = BytesIO()





    workbook.save(output)





    output.seek(0)





    return StreamingResponse(



        output,



        media_type=(



            "application/"



            "vnd.openxmlformats-officedocument."



            "spreadsheetml.sheet"



        ),



        headers={



            "Content-Disposition":



                (



                    "attachment; "



                    "filename="



                    "Template_Master_Lembaga.xlsx"



                )



        },



    )





# ============================================================



# EXPORT EXCEL



# ============================================================





@app.get(



    "/institutions/export/xlsx"



)



def export_institutions_xlsx():





    rows = get_institution_export_data()





    workbook = Workbook()





    worksheet = workbook.active





    worksheet.title = (



        "Master Data Lembaga"



    )





    for (



        column_index,



        column_name,



    ) in enumerate(



        INSTITUTION_TEMPLATE_COLUMNS,



        start=1,



    ):





        cell = worksheet.cell(



            row=1,



            column=column_index,



            value=column_name,



        )





        cell.font = Font(



            bold=True



        )





        cell.alignment = Alignment(



            horizontal="center",



            vertical="center",



        )





    for row_index, item in enumerate(



        rows,



        start=2,



    ):





        for (



            column_index,



            column_name,



        ) in enumerate(



            INSTITUTION_TEMPLATE_COLUMNS,



            start=1,



        ):





            value = (



                item[



                    column_name



                ]



            )





            cell = worksheet.cell(



                row=row_index,



                column=column_index,



                value=(



                    value



                    if value is not None



                    else ""



                ),



            )





            cell.alignment = Alignment(



                vertical="top",



                wrap_text=True,



            )





    worksheet.freeze_panes = "A2"





    worksheet.auto_filter.ref = (



        worksheet.dimensions



    )





    widths = {



        "A": 20,



        "B": 40,



        "C": 24,



        "D": 24,



        "E": 24,



        "F": 45,



        "G": 50,



        "H": 25,



        "I": 25,



        "J": 15,



        "K": 25,



        "L": 25,



        "M": 40,



        "N": 30,



        "O": 25,



    }





    for (



        column_letter,



        width,



    ) in widths.items():





        worksheet.column_dimensions[



            column_letter



        ].width = width





    output = BytesIO()





    workbook.save(output)





    output.seek(0)





    return StreamingResponse(



        output,



        media_type=(



            "application/"



            "vnd.openxmlformats-officedocument."



            "spreadsheetml.sheet"



        ),



        headers={



            "Content-Disposition":



                (



                    "attachment; "



                    "filename="



                    "Master_Data_Lembaga.xlsx"



                )



        },



    )





# ============================================================



# EXPORT CSV



# ============================================================





@app.get(



    "/institutions/export/csv"



)



def export_institutions_csv():





    rows = get_institution_export_data()





    output = StringIO(



        newline=""



    )





    writer = csv.DictWriter(



        output,



        fieldnames=(



            INSTITUTION_TEMPLATE_COLUMNS



        ),



    )





    writer.writeheader()





    for item in rows:





        writer.writerow(



            {



                column:



                    (



                        item[column]





                        if item[column]



                        is not None





                        else ""



                    )





                for column



                in INSTITUTION_TEMPLATE_COLUMNS



            }



        )





    csv_content = (



        "\ufeff"



        + output.getvalue()



    )





    csv_bytes = (



        csv_content.encode(



            "utf-8"



        )



    )





    return StreamingResponse(



        iter(



            [



                csv_bytes



            ]



        ),



        media_type=(



            "text/csv; charset=utf-8"



        ),



        headers={



            "Content-Disposition":



                (



                    "attachment; "



                    "filename="



                    "Master_Data_Lembaga.csv"



                )



        },



    )





# ============================================================



# MASTER LEMBAGA LIST



# ============================================================





@app.get(



    "/institutions",



    response_class=HTMLResponse,



)



def institution_list(



    request: Request,



    search: str = "",



    status: str = "",



    result: str = "",



    message: str = "",



):





    conn = get_connection()





    try:





        query = """



            SELECT





                institution.*,





                COUNT(



                    institution_amendment.amendment_id



                ) AS amendment_count





            FROM institution





            LEFT JOIN institution_amendment





                ON institution.institution_id



                =



                institution_amendment.institution_id





            WHERE institution.source_status =



                'AVAILABLE'



        """





        params = []





        if search.strip():





            query += """



                AND (



                    LOWER(



                        institution.id_lspro



                    ) LIKE LOWER(?)





                    OR LOWER(



                        institution.name_lspro



                    ) LIKE LOWER(?)





                    OR LOWER(



                        institution.kab_kota



                    ) LIKE LOWER(?)





                    OR LOWER(



                        institution.propinsi



                    ) LIKE LOWER(?)





                    OR LOWER(



                        institution.cp



                    ) LIKE LOWER(?)



                )



            """





            keyword = (



                f"%{search.strip()}%"



            )





            params.extend(



                [



                    keyword,



                    keyword,



                    keyword,



                    keyword,



                    keyword,



                ]



            )





        if status.strip():





            query += """



                AND LOWER(



                    institution.status_accreditation



                ) = LOWER(?)



            """





            params.append(



                status.strip()



            )





        query += """



            GROUP BY



                institution.institution_id





            ORDER BY



                institution.id_lspro



        """





        institutions = (



            conn.execute(



                query,



                params,



            )



            .fetchall()



        )





    finally:





        conn.close()





    return templates.TemplateResponse(



        request=request,



        name="institutions/list.html",



        context={



            "institutions":



                institutions,





            "search":



                search,





            "status":



                status,





            "result":



                result,





            "message":



                message,



        },



    )





# ============================================================



# EDIT PAGE



# ============================================================





@app.get(



    "/institutions/{institution_id}/edit",



    response_class=HTMLResponse,



)



def institution_edit_page(



    request: Request,



    institution_id: int,



    status: str = "",



    message: str = "",



):





    conn = get_connection()





    try:





        institution = (



            conn.execute(



                """



                SELECT *





                FROM institution





                WHERE institution_id = ?



                """,



                (



                    institution_id,



                ),



            )



            .fetchone()



        )





        if institution is None:





            raise HTTPException(



                status_code=404,



                detail=(



                    "Data lembaga tidak ditemukan."



                ),



            )





        amendments = (



            conn.execute(



                """



                SELECT



                    amendment_id,



                    amendment_no,



                    amendment_date





                FROM institution_amendment





                WHERE institution_id = ?





                ORDER BY amendment_no ASC



                """,



                (



                    institution_id,



                ),



            )



            .fetchall()



        )





    finally:





        conn.close()





    return templates.TemplateResponse(



        request=request,



        name="institutions/edit.html",



        context={



            "institution":



                institution,





            "amendments":



                amendments,





            "status":



                status,





            "message":



                message,



        },



    )





# ============================================================



# EDIT PROCESS



# ============================================================





@app.post(



    "/institutions/{institution_id}/edit"



)



def institution_edit_process(





    institution_id: int,





    id_lspro: str = Form(...),





    name_lspro: str = Form(...),





    date_start_accreditation: str = Form(""),





    date_end_accreditation: str = Form(""),





    status_accreditation: str = Form(""),





    address: str = Form(""),





    kab_kota: str = Form(""),





    propinsi: str = Form(""),





    kode_pos: str = Form(""),





    telp: str = Form(""),





    fax: str = Form(""),





    email: str = Form(""),





    cp: str = Form(""),





    no_hp_cp: str = Form(""),





    amendment_no: list[str] = Form([]),





    amendment_date: list[str] = Form([]),



):





    conn = get_connection()





    try:





        existing = (



            conn.execute(



                """



                SELECT *





                FROM institution





                WHERE institution_id = ?



                """,



                (



                    institution_id,



                ),



            )



            .fetchone()



        )





        if existing is None:





            raise HTTPException(



                status_code=404,



                detail=(



                    "Data lembaga tidak ditemukan."



                ),



            )





        if not id_lspro.strip():





            raise ValueError(



                "ID LSPro wajib diisi."



            )





        if not name_lspro.strip():





            raise ValueError(



                "Nama lembaga wajib diisi."



            )





        duplicate = (



            conn.execute(



                """



                SELECT institution_id





                FROM institution





                WHERE id_lspro = ?





                AND institution_id != ?



                """,



                (



                    id_lspro,



                    institution_id,



                ),



            )



            .fetchone()



        )





        if duplicate is not None:





            raise ValueError(



                f"ID LSPro {id_lspro} "



                f"sudah digunakan oleh lembaga lain."



            )





        amendment_rows = []





        used_numbers = set()





        total_amendments = max(



            len(amendment_no),



            len(amendment_date),



        )





        for index in range(



            total_amendments



        ):





            no_value = (



                amendment_no[index]



                if index < len(amendment_no)



                else ""



            )





            date_value = (



                amendment_date[index]



                if index < len(amendment_date)



                else ""



            )





            no_value = no_value.strip()



            date_value = date_value.strip()





            if (



                not no_value



                and



                not date_value



            ):





                continue





            if not no_value:





                raise ValueError(



                    "Nomor amandemen wajib diisi "



                    "jika tanggal amandemen diisi."



                )





            try:





                number = int(no_value)





            except ValueError:





                raise ValueError(



                    f"Nomor amandemen '{no_value}' "



                    f"harus berupa angka."



                )





            if number <= 0:





                raise ValueError(



                    "Nomor amandemen harus lebih dari 0."



                )





            if number in used_numbers:





                raise ValueError(



                    f"Amandemen ke-{number} duplikat."



                )





            used_numbers.add(number)





            amendment_rows.append(



                {



                    "amendment_no":



                        number,





                    "amendment_date":



                        (



                            date_value



                            or None



                        ),



                }



            )





        conn.execute(



            "BEGIN"



        )





        conn.execute(



            """



            UPDATE institution





            SET



                id_lspro = ?,



                name_lspro = ?,



                date_start_accreditation = ?,



                date_end_accreditation = ?,



                status_accreditation = ?,



                address = ?,



                kab_kota = ?,



                propinsi = ?,



                kode_pos = ?,



                telp = ?,



                fax = ?,



                email = ?,



                cp = ?,



                no_hp_cp = ?,



                updated_at = CURRENT_TIMESTAMP





            WHERE institution_id = ?



            """,



            (



                id_lspro,





                name_lspro,





                date_start_accreditation



                or None,





                date_end_accreditation



                or None,





                status_accreditation



                or None,





                address



                or None,





                kab_kota



                or None,





                propinsi



                or None,





                kode_pos



                or None,





                telp



                or None,





                fax



                or None,





                email



                or None,





                cp



                or None,





                no_hp_cp



                or None,





                institution_id,



            ),



        )





        conn.execute(



            """



            DELETE FROM institution_amendment





            WHERE institution_id = ?



            """,



            (



                institution_id,



            ),



        )





        for item in amendment_rows:





            conn.execute(



                """



                INSERT INTO institution_amendment (



                    institution_id,



                    amendment_no,



                    amendment_date



                )





                VALUES (



                    ?,



                    ?,



                    ?



                )



                """,



                (



                    institution_id,





                    item[



                        "amendment_no"



                    ],





                    item[



                        "amendment_date"



                    ],



                ),



            )





        conn.commit()





    except HTTPException:





        conn.rollback()



        raise





    except Exception as exc:





        conn.rollback()





        query = urlencode(



            {



                "status":



                    "failed",





                "message":



                    str(exc),



            }



        )





        return RedirectResponse(



            url=(



                f"/institutions/"



                f"{institution_id}/edit?"



                + query



            ),



            status_code=303,



        )





    finally:





        conn.close()





    query = urlencode(



        {



            "result":



                "success",





            "message":



                "Data lembaga berhasil diperbarui.",



        }



    )





    return RedirectResponse(



        url=(



            "/institutions?"



            + query



        ),



        status_code=303,



    )





# ============================================================



# DETAIL LEMBAGA



# ============================================================





@app.get(



    "/institutions/{institution_id}",



    response_class=HTMLResponse,



)



def institution_detail(



    request: Request,



    institution_id: int,



):





    conn = get_connection()





    try:





        institution = (



            conn.execute(



                """



                SELECT *





                FROM institution





                WHERE institution_id = ?



                """,



                (



                    institution_id,



                ),



            )



            .fetchone()



        )





        if institution is None:





            raise HTTPException(



                status_code=404,



                detail=(



                    "Data lembaga tidak ditemukan."



                ),



            )





        amendments = (



            conn.execute(



                """



                SELECT *





                FROM institution_amendment





                WHERE institution_id = ?





                ORDER BY amendment_no ASC



                """,



                (



                    institution_id,



                ),



            )



            .fetchall()



        )





        amendment_count = (



            len(amendments)



        )





        latest_amendment = (



            amendments[-1]





            if amendments





            else None



        )





    finally:





        conn.close()





    return templates.TemplateResponse(



        request=request,



        name="institutions/detail.html",



        context={



            "institution":



                institution,





            "amendments":



                amendments,





            "amendment_count":



                amendment_count,





            "latest_amendment":



                latest_amendment,



        },



    )
